from PyQt5.QtWidgets import (QWidget, QVBoxLayout, QPushButton,
                             QLabel, QFileDialog, QMessageBox, QTableWidget,
                             QTableWidgetItem, QSizePolicy)
from PyQt5.QtCore import Qt, pyqtSignal
from PyQt5.QtGui import QPixmap
import openpyxl
from openpyxl.chart import PieChart, Reference
from openpyxl.utils import get_column_letter
import matplotlib.pyplot as plt
from io import BytesIO
import tempfile
import os
from windows.mixin import AutoRefreshMixin


class ReportWindow(QWidget, AutoRefreshMixin):
    closed = pyqtSignal()

    def __init__(self, controller):
        QWidget.__init__(self)
        AutoRefreshMixin.__init__(self, controller)
        self.ctrl = controller
        self.setWindowTitle("Отчёт")
        self._build_ui()
        self._refresh_preview()
        self.refresh()

    def _build_ui(self):
        label = QLabel("Предпросмотр отчёта")
        label.setStyleSheet("font-size:16pt; font-weight:600;")

        self.table = QTableWidget()
        self.table.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.pic_label = QLabel("Круговая диаграмма появится после расчёта")
        self.pic_label.setAlignment(Qt.AlignCenter)
        self.pic_label.setMinimumHeight(300)

        self.btn = QPushButton("Сохранить отчет")
        back_btn = QPushButton("Назад")
        self.btn.clicked.connect(self._save_excel)
        back_btn.clicked.connect(self.closed.emit)

        lo = QVBoxLayout(self)
        lo.addWidget(label)
        lo.addWidget(self.table, stretch=2)
        lo.addWidget(self.pic_label, stretch=1)
        lo.addWidget(self.btn)
        lo.addWidget(back_btn)

    def _refresh_preview(self):
        data, total_qty, total_cost = self.ctrl.report_data()
        self.table.setColumnCount(7)
        self.table.setHorizontalHeaderLabels(
            ["ID", "Название", "Категория", "Кол-во", "Цена", "Сумма", "Добавлено"])
        self.table.setRowCount(len(data) + 2)
        for r, p in enumerate(data):
            for c, key in enumerate(
                    ["id", "name", "category", "quantity", "price"]):
                self.table.setItem(r, c, QTableWidgetItem(str(p[key])))
            self.table.setItem(r, 5, QTableWidgetItem(str(p['quantity'] * p['price'])))
            self.table.setItem(r, 6, QTableWidgetItem(p['date_added']))
        self.table.setItem(len(data), 0, QTableWidgetItem("Итого"))
        self.table.setItem(len(data), 3, QTableWidgetItem(str(total_qty)))
        self.table.setItem(len(data), 5, QTableWidgetItem(str(total_cost)))

        self._build_pie_chart(data)

    def _build_pie_chart(self, data):
        cat_qty = {}
        for p in data:
            cat = p['category']
            cat_qty[cat] = cat_qty.get(cat, 0) + p['quantity']
        if not cat_qty:
            self.pic_label.setText("Нет данных для диаграммы")
            return
        labels = list(cat_qty.keys())
        sizes  = list(cat_qty.values())

        plt.ioff()
        fig, ax = plt.subplots(figsize=(3, 3))
        ax.pie(sizes, labels=labels, autopct='%1.1f%%', startangle=90,
               colors=plt.cm.Set3.colors)
        ax.axis('equal')
        buffer = BytesIO()
        fig.savefig(buffer, format='png', bbox_inches='tight')
        plt.close(fig)
        buffer.seek(0)
        pixmap = QPixmap()
        pixmap.loadFromData(buffer.read())
        self.pic_label.setPixmap(pixmap)

    def _save_excel(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Сохранить отчёт", "report.xlsx",
            "Excel files (*.xlsx)")
        if not path:
            return
        try:
            data, total_qty, total_cost = self.ctrl.report_data()
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Склад"

            headers = ["ID", "Название", "Категория", "Кол-во", "Цена", "Сумма", "Добавлено"]
            ws.append(headers)

            for p in data:
                ws.append([p['id'], p['name'], p['category'],
                           p['quantity'], p['price'],
                           p['quantity'] * p['price'], p['date_added']])

            ws.append([])
            ws.append(["Итого", "", "", total_qty, "", total_cost])

            cat_qty = {}
            for p in data:
                cat = p['category']
                cat_qty[cat] = cat_qty.get(cat, 0) + p['quantity']
            chart = PieChart()
            labels = Reference(ws, min_col=3, min_row=2,
                               max_row=1 + len(cat_qty))
            values = Reference(ws, min_col=4, min_row=2,
                               max_row=1 + len(cat_qty))
            chart.add_data(values, titles_from_data=False)
            chart.set_categories(labels)
            chart.title = "Распределение по категориям (шт.)"
            ws.add_chart(chart, "I2")

            wb.save(path)
            QMessageBox.information(self, "Готово", f"Отчёт сохранён:\n{path}")
        except Exception as e:
            QMessageBox.warning(self, "Ошибка", str(e))

    ...
    def refresh(self):
        data, total_qty, total_cost = self.ctrl.report_data()
        self.table.setRowCount(len(data) + 1)
        for r, p in enumerate(data):
            for c, key in enumerate(
                    ["id", "name", "category", "quantity", "price"]):
                self.table.setItem(r, c, QTableWidgetItem(str(p[key])))
            self.table.setItem(r, 5, QTableWidgetItem(str(p['quantity'] * p['price'])))
            self.table.setItem(r, 6, QTableWidgetItem(p['date_added']))
        last = len(data)
        self.table.setItem(last, 0, QTableWidgetItem("Итого"))
        self.table.setItem(last, 3, QTableWidgetItem(str(total_qty)))
        self.table.setItem(last, 5, QTableWidgetItem(str(total_cost)))
        self._build_pie_chart(data)